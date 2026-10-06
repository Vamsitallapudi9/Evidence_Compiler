import io
import json

import docx
import httpx
import pytest

from evidence_compiler.core.models import Creator, EvidenceOrigin, SourceKind
from evidence_compiler.ingestion.documents import ingest_docx, parse_html
from evidence_compiler.ingestion.fetch import FetchError, check_host, fetch, rewrite
from evidence_compiler.ingestion.url import UrlIngestor, find_urls

PUBLIC = "http://93.184.216.34"  # an IP literal: no DNS needed, and it is public

HTML = b"""<html><head><title>Quarterly update</title><script>var x="ignore";</script></head>
<body><nav>Home | About | Contact us today</nav><main><h1>Results for the quarter</h1>
<p>Revenue for the quarter was $114.0 million, a decrease of 8.4% year over year.</p>
<p>Europe was the weakest region, with several enterprise customers not renewing.</p>
</main><footer>Copyright 2026 Example Corp. All rights reserved.</footer></body></html>"""


def transport(routes: dict[str, httpx.Response]) -> httpx.MockTransport:
    def handler(req: httpx.Request) -> httpx.Response:
        return routes.get(str(req.url), httpx.Response(404))

    return httpx.MockTransport(handler)


@pytest.mark.parametrize(
    ("given", "expected"),
    [
        (
            "https://docs.google.com/document/d/abc123/edit?usp=sharing",
            "https://docs.google.com/document/d/abc123/export?format=docx",
        ),
        (
            "https://docs.google.com/spreadsheets/d/xyz/edit#gid=42",
            "https://docs.google.com/spreadsheets/d/xyz/export?format=csv&gid=42",
        ),
        (
            "https://drive.google.com/file/d/FILE/view",
            "https://drive.google.com/uc?export=download&id=FILE",
        ),
        (
            "https://github.com/o/r/blob/main/README.md",
            "https://raw.githubusercontent.com/o/r/main/README.md",
        ),
        ("https://arxiv.org/abs/2401.01234", "https://arxiv.org/pdf/2401.01234"),
        ("https://example.com/a.pdf", "https://example.com/a.pdf"),
    ],
)
def test_share_links_are_rewritten(given, expected):
    assert rewrite(given) == expected


@pytest.mark.parametrize(
    "url",
    [
        "file:///etc/passwd",
        "ftp://example.com/x",
        "http://127.0.0.1:8000/api/sources",
        "http://localhost/",
        "http://10.0.0.5/",
        "http://192.168.1.1/",
        "http://169.254.169.254/latest/meta-data/",
        "http://[::1]/",
        "http://user:pass@example.com/",
    ],
)
def test_ssrf_targets_are_refused(url):
    with pytest.raises(FetchError):
        check_host(url)


def test_redirect_to_private_address_is_refused():
    t = transport(
        {f"{PUBLIC}/r": httpx.Response(302, headers={"location": "http://127.0.0.1/secret"})}
    )
    with pytest.raises(FetchError, match="non-public"):
        fetch(f"{PUBLIC}/r", transport=t)


def test_size_limit_and_http_errors():
    t = transport(
        {
            f"{PUBLIC}/big": httpx.Response(200, content=b"x" * 2048),
            f"{PUBLIC}/private": httpx.Response(403),
        }
    )
    with pytest.raises(FetchError, match="limit"):
        fetch(f"{PUBLIC}/big", max_bytes=1024, transport=t)
    with pytest.raises(FetchError, match="sign-in") as e:
        fetch(f"{PUBLIC}/private", transport=t)
    assert e.value.status == 403


def test_html_extraction_skips_chrome():
    title, blocks = parse_html(HTML)
    assert title == "Quarterly update"
    assert any("8.4%" in b for b in blocks)
    joined = " ".join(blocks)
    assert "Contact us" not in joined and "Copyright" not in joined and "var x" not in joined


def test_docx_sections_and_tables():
    d = docx.Document()
    d.add_heading("Summary", level=1)
    d.add_paragraph("Revenue declined 8.4% in the second quarter.")
    d.add_heading("Regions", level=1)
    d.add_paragraph("Europe declined the most.")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Region", "Revenue"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Europe", "$34.2M"
    buf = io.BytesIO()
    d.save(buf)
    result = ingest_docx(buf.getvalue(), name="r.docx", uri="r.docx")
    assert result.source.kind == SourceKind.DOCX
    assert len(result.pages) == 3  # two heading sections + the table
    for e in result.evidence:
        assert e.origin == EvidenceOrigin.PRIMARY and e.creator == Creator.CONNECTOR
        assert result.pages[e.location.page][e.location.char_start : e.location.char_end] == e.text
    assert any(e.text == "Region: Europe | Revenue: $34.2M" for e in result.evidence)


def test_find_urls():
    q = "Summarise https://example.com/a.pdf, and compare with (https://x.org/page)."
    assert find_urls(q) == ["https://example.com/a.pdf", "https://x.org/page"]


@pytest.fixture
def link_ctx(ctx, demo_dir):
    pdf = (demo_dir / "q2_fy2026_report.pdf").read_bytes()
    csv = (demo_dir / "regional_sales.csv").read_bytes()
    ctx.urls = UrlIngestor(
        ctx.settings.uploads_dir,
        ctx.registry.tabular,
        max_bytes=10**7,
        transport=transport(
            {
                f"{PUBLIC}/report": httpx.Response(
                    200, content=pdf, headers={"content-type": "application/pdf"}
                ),
                f"{PUBLIC}/page": httpx.Response(
                    200, content=HTML, headers={"content-type": "text/html"}
                ),
                f"{PUBLIC}/data.csv": httpx.Response(
                    200, content=csv, headers={"content-type": "text/csv"}
                ),
                f"{PUBLIC}/bin": httpx.Response(
                    200, content=b"\x00\x01", headers={"content-type": "application/zip"}
                ),
            }
        ),
    )
    return ctx


def test_ingest_links_of_each_kind(link_ctx):
    pdf = link_ctx.ingest_url(f"{PUBLIC}/report")
    assert pdf.kind == SourceKind.PDF and pdf.uri == f"{PUBLIC}/report"
    assert pdf.metadata["url"] == f"{PUBLIC}/report" and "fetched_at" in pdf.metadata
    web = link_ctx.ingest_url(f"{PUBLIC}/page")
    assert web.kind == SourceKind.WEB and web.name == "Quarterly update"
    items = [e for e in link_ctx.repo.all_evidence() if web.id in e.source_ids]
    page, _ = link_ctx.repo.page_text(web.id, 1)
    assert items and all(page[e.location.char_start : e.location.char_end] == e.text for e in items)
    csv = link_ctx.ingest_url(f"{PUBLIC}/data.csv")
    assert csv.kind == SourceKind.CSV and csv.metadata["tables"] == ["data"]
    with pytest.raises(FetchError, match="unsupported"):
        link_ctx.ingest_url(f"{PUBLIC}/bin")


def test_run_with_link_is_scoped_to_that_link(link_ctx, monkeypatch):
    seen = {}

    class Stop(Exception):
        pass

    def fake_run(self, question, *, run_id, emit=None, source_ids=None):
        seen["scope"] = source_ids
        raise Stop

    monkeypatch.setattr("evidence_compiler.pipeline.orchestrator.Pipeline.run", fake_run)
    from evidence_compiler.app import _LiveRun
    from evidence_compiler.llm.meter import MeteredLLM
    from tests.fake_scenario import make_fake

    run_id = "run_x"
    link_ctx.repo.create_run(run_id, "q")
    link_ctx.runs._live[run_id] = _LiveRun(meter=MeteredLLM(make_fake(), link_ctx.settings))
    link_ctx.runs._work(run_id, f"What does {PUBLIC}/page say about Europe?")
    web = next(s for s in link_ctx.repo.list_sources() if s.kind == SourceKind.WEB)
    assert seen["scope"] == {web.id}
    stages = [e.stage for e in link_ctx.runs.events(run_id)]
    assert stages[:2] == ["fetch", "fetch"]


def test_scoped_pipeline_ignores_other_sources(ctx):
    from evidence_compiler.pipeline.orchestrator import Pipeline
    from tests.fake_scenario import make_fake

    q1 = next(s for s in ctx.repo.list_sources() if s.name == "q1_fy2026_report.pdf")
    keep = {s.id for s in ctx.repo.list_sources() if s.id != q1.id}
    pipe = Pipeline(make_fake(), ctx.repo, ctx.retriever, ctx.executor, ctx.settings)
    answer = pipe.run("Why did Q2 revenue decline?", run_id="r", source_ids=keep)
    assert q1.id not in {sid for e in answer.evidence.values() for sid in e.source_ids}


def test_huggingface_connector(ctx):
    from evidence_compiler.ingestion.huggingface import HuggingFaceConnector, normalize_name

    assert (
        normalize_name("https://huggingface.co/datasets/stanfordnlp/imdb/tree/main")
        == "stanfordnlp/imdb"
    )
    with pytest.raises(FetchError):
        normalize_name("not a name; drop table")

    review = "This film was a remarkable achievement in storytelling and acting, " * 3
    rows = [
        {"row_idx": i, "row": {"text": f"{review} #{i}", "label": i % 2, "meta": {"k": i}}}
        for i in range(150)
    ]
    features = [
        {"name": "text", "type": {"_type": "Value", "dtype": "string"}},
        {"name": "label", "type": {"_type": "ClassLabel", "names": ["neg", "pos"]}},
        {"name": "meta", "type": {"k": {"_type": "Value"}}},
    ]

    def handler(req: httpx.Request) -> httpx.Response:
        url = str(req.url)
        if "/api/datasets/" in url:
            return httpx.Response(200, json={"cardData": {"license": "apache-2.0"}})
        if "/splits" in url:
            return httpx.Response(
                200,
                json={
                    "splits": [
                        {"dataset": "o/reviews", "config": "plain_text", "split": "test"},
                        {"dataset": "o/reviews", "config": "plain_text", "split": "train"},
                    ]
                },
            )
        if "/rows" in url:
            off, n = int(req.url.params["offset"]), int(req.url.params["length"])
            assert req.url.params["split"] == "train"
            return httpx.Response(
                200,
                json={
                    "features": features,
                    "rows": rows[off : off + n],
                    "num_rows_total": len(rows),
                },
            )
        return httpx.Response(404)

    conn = HuggingFaceConnector(ctx.tabular, transport=httpx.MockTransport(handler))
    result = conn.register("o/reviews", max_rows=120)
    ctx._store(result)
    src = result.source
    assert src.kind == SourceKind.HUGGINGFACE and src.metadata["license"] == "apache-2.0"
    assert src.metadata["rows_loaded"] == 120 and src.metadata["text_columns"] == ["text"]
    assert len(result.evidence) == 120 and result.evidence[0].metadata["passage"]
    assert src.metadata["sampling"] == "spread"
    row_ids = sorted(e.location.row_index for e in result.evidence)
    assert row_ids[0] == 0 and row_ids[-1] > 100  # spans the split, original row numbers
    q = ctx.executor.run(
        "tabular",
        "SELECT label_name, COUNT(*) AS n FROM hf_reviews_train GROUP BY label_name ORDER BY 1",
        ctx.repo.schema_cards(),
    )
    assert [(e.fields["label_name"], e.fields["n"]) for e in q.evidence] == [
        ("neg", 60),
        ("pos", 60),
    ]
    assert json.loads(
        ctx.tabular.query("SELECT meta FROM hf_reviews_train LIMIT 1", 1)[1][0][0]
    ) == {"k": 0}
