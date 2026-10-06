"""Text-document connectors: web pages (HTML), Word documents and plain text.

Like the PDF connector, each produces one PRIMARY evidence item per paragraph, with a
character span into a canonical per-section text, so passages can be highlighted.
"""

import io
import re
from html.parser import HTMLParser

from evidence_compiler.core.ids import content_id, file_hash
from evidence_compiler.core.models import (
    Creator,
    EvidenceItem,
    EvidenceOrigin,
    Source,
    SourceKind,
    TextLocation,
)
from evidence_compiler.ingestion.base import IngestionError, IngestResult, check_zip_container
from evidence_compiler.security.untrusted import scan_for_injection

_MAX_PARAGRAPH = 1500  # split very long blocks so citations stay precise


def _split_long(text: str) -> list[str]:
    if len(text) <= _MAX_PARAGRAPH:
        return [text]
    sentences = re.split(r"(?<=[.!?])\s+", text)
    out, cur = [], ""
    for s in sentences:
        if cur and len(cur) + len(s) + 1 > _MAX_PARAGRAPH:
            out.append(cur)
            cur = s
        else:
            cur = f"{cur} {s}".strip()
    if cur:
        out.append(cur)
    return out


def build_text_source(
    *,
    kind: SourceKind,
    name: str,
    uri: str,
    data: bytes,
    sections: list[list[str]],
    connector: str,
    metadata: dict | None = None,
) -> IngestResult:
    source_id = content_id("src", kind.value, file_hash(data))
    source = Source(
        id=source_id,
        kind=kind,
        name=name,
        uri=uri,
        content_hash=file_hash(data),
        metadata=metadata or {},
    )
    result = IngestResult(source=source)
    for n, paragraphs in enumerate((s for s in sections if s), start=1):
        text = ""
        for para in (p for raw in paragraphs for p in _split_long(" ".join(raw.split()))):
            if not para:
                continue
            if text:
                text += "\n\n"
            start = len(text)
            text += para
            loc = TextLocation(page=n, char_start=start, char_end=len(text))
            result.evidence.append(
                EvidenceItem(
                    id=content_id("ev", source_id, loc.model_dump(), para),
                    source_ids=[source_id],
                    origin=EvidenceOrigin.PRIMARY,
                    creator=Creator.CONNECTOR,
                    text=para,
                    location=loc,
                    security_flags=scan_for_injection(para),
                    metadata={"connector": connector},
                )
            )
        result.pages[n] = text
    if not result.evidence:
        raise IngestionError("no readable text found in the document")
    source.metadata["page_count"] = len(result.pages)
    return result


# --- HTML --------------------------------------------------------------------------------

_SKIP = {
    "script",
    "style",
    "noscript",
    "template",
    "svg",
    "nav",
    "footer",
    "aside",
    "form",
    "button",
    "iframe",
    "select",
    "header",
}
_BLOCK = {
    "p",
    "div",
    "section",
    "article",
    "main",
    "li",
    "td",
    "th",
    "tr",
    "pre",
    "blockquote",
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "dd",
    "dt",
    "figcaption",
    "caption",
    "br",
    "table",
    "ul",
    "ol",
}


class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[str] = []
        self.title = ""
        self._buf: list[str] = []
        self._skip = 0
        self._in_title = False
        self._has_main = False

    def handle_starttag(self, tag: str, attrs: list) -> None:
        if tag in _SKIP:
            self._skip += 1
        elif tag == "title":
            self._in_title = True
        elif tag in _BLOCK:
            self._flush()

    def handle_endtag(self, tag: str) -> None:
        if tag in _SKIP:
            self._skip = max(0, self._skip - 1)
        elif tag == "title":
            self._in_title = False
        elif tag in _BLOCK:
            self._flush()

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data
        elif not self._skip:
            self._buf.append(data)

    def _flush(self) -> None:
        text = " ".join("".join(self._buf).split())
        self._buf = []
        if len(text) >= 3:
            self.blocks.append(text)

    def close(self) -> None:
        super().close()
        self._flush()


def parse_html(data: bytes) -> tuple[str, list[str]]:
    html = data.decode("utf-8", errors="replace")
    m = re.search(r"<(article|main)\b.*?</\1>", html, re.S | re.I)
    parser = _TextExtractor()
    parser.feed(m.group(0) if m and len(m.group(0)) > 2000 else html)
    parser.close()
    title_parser = _TextExtractor()
    title_parser.feed(html[:20000])
    title = " ".join((title_parser.title or parser.title).split())
    # Drop very short UI fragments ("Share", "Menu") but keep headings with some words.
    blocks = [b for b in parser.blocks if len(b) >= 25 or len(b.split()) >= 3]
    return title, blocks


def ingest_html(data: bytes, *, name: str | None, url: str, metadata: dict) -> IngestResult:
    title, blocks = parse_html(data)
    if not blocks:
        raise IngestionError("no readable text on that page (it may require JavaScript or sign-in)")
    return build_text_source(
        kind=SourceKind.WEB,
        name=name or title or url,
        uri=url,
        data=data,
        sections=[blocks],
        connector="web",
        metadata={**metadata, "title": title},
    )


# --- Word / text -------------------------------------------------------------------------


def ingest_docx(data: bytes, *, name: str, uri: str, metadata: dict | None = None) -> IngestResult:
    import docx  # python-docx

    check_zip_container(data)
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types
        raise IngestionError(f"could not read Word document: {exc}") from exc
    sections: list[list[str]] = [[]]
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading 1") and sections[-1]:
            sections.append([])  # each top-level heading starts a new section
        sections[-1].append(text)
    for table in document.tables:
        header = [c.text.strip() for c in table.rows[0].cells] if table.rows else []
        rows = []
        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            rows.append(" | ".join(f"{h}: {v}" for h, v in zip(header, cells, strict=False)))
        if rows:
            sections.append(rows)
    return build_text_source(
        kind=SourceKind.DOCX,
        name=name,
        uri=uri,
        data=data,
        sections=sections,
        connector="docx",
        metadata=metadata,
    )


def ingest_text(data: bytes, *, name: str, uri: str, metadata: dict | None = None) -> IngestResult:
    text = data.decode("utf-8", errors="replace")
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    return build_text_source(
        kind=SourceKind.TEXT,
        name=name,
        uri=uri,
        data=data,
        sections=[paragraphs],
        connector="text",
        metadata=metadata,
    )


class DocxConnector:
    def ingest(self, path, *, name: str) -> IngestResult:
        return ingest_docx(path.read_bytes(), name=name, uri=str(path))


class TextConnector:
    def ingest(self, path, *, name: str) -> IngestResult:
        return ingest_text(path.read_bytes(), name=name, uri=str(path))
